"""One-off deterministic seed script for the stakeholders' master catalog,
backend/evals/master-product-detail-oct-26.xlsx (the "Master Product Detail
File (Oct-26) - Do Not Edit" workbook, committed as-is and read directly),
into the Neo4j graph (see backend/retrieval/graph_client.py's schema
docstring) and its Postgres mirror (backend/catalog/repository.py). Run
manually — this is NOT part of per-document ingestion
(backend/ingestion/pipeline.py); it's a batch job seeding the authoritative
ProductFamily > Tray > Part catalog that per-document prose extraction
(backend/ingestion/entity_extraction.py) then attaches facts to. Run it (and
backend/ingestion/seed_synonyms.py) before uploading/indexing real
documents, or prose extraction has nothing to attach to:

    python -m ingestion.seed_master_catalog [path/to/master.xlsx]

Columns are looked up by header name, never position. `Item No.` is a
Part's only identity — rows without one are skipped. `Family` is strict:
every row needs one (the seed fails loudly otherwise) and every Tray must
sit under exactly one Family. Tray names drop the leading
`UNITE® Foot & Ankle | ` brand segment every row repeats.

Two kinds of relationship are resolved deterministically here, both
acceptable only inside this one-off, reviewable script (never in the LLM
prose-extraction path, which only merges by exact SKU):

  - COMPATIBLE_WITH (plate <-> screw family): each compatibility-matrix
    column header names a literal SKU prefix, e.g. "2.7mm Polyaxial Locking
    (MPSL27xx)" means any screw whose Item No. starts with "MPSL27" is
    compatible with a plate row marked "X" in that column. Exact prefix
    match, no fuzzy text involved.
  - REQUIRES_TOOL (screw <-> guidewire/drill-bit/driver): the Guidewire/
    Pre-Drill Diameter/Driver columns are free text with no SKU reference,
    even though matching catalog rows exist. Resolved by matching numeric
    tokens (diameter/length) or driver-bit tokens against candidate rows'
    Description, restricted to the same tray to avoid cross-tray false
    positives, and only written when exactly one candidate matches —
    ambiguous cases are left as the flat spec property on the Part instead
    of guessed into an edge.

Tray membership is reconciled, not just merged: after upserting, every
listed Part's BELONGS_TO_TRAY edges are trimmed to exactly the trays this
file gives it, and a Tray left empty is deleted unless something else (a
Procedure from document extraction) still points at it — so renaming or
moving a tray in the master file needs no hand-run Cypher cleanup. Parts
dropped from the file entirely are left alone; deleting them would also
delete facts other ingestion paths attached to them.
"""
import asyncio
import re
import sys
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from catalog.repository import upsert_families_and_trays
from retrieval.graph_client import GraphClient, get_graph_client

DEFAULT_PATH = (
    Path(__file__).resolve().parents[2] / "backend" / "evals" / "master-product-detail-oct-26.xlsx"
)

_BRAND_PREFIX = "UNITE® Foot & Ankle | "
_FIXED_COLUMNS = {
    "Family",
    "Tray",
    "Item Type",
    "Item No.",
    "Description",
    "Qty per Set",
    "Head Style",
    "Construct",
    "Thread",
    "Color",
    "Guidewire",
    "Pre-Drill Diameter",
    "Driver",
}
_REQUIRED_COLUMNS = ("Family", "Tray", "Item No.")
# The closing paren is optional: several Oct-26 headers lost theirs
# ("2.0mm Non-Locking (MMSN20xx"), and the prefix is still unambiguous.
_COMPAT_PREFIX_RE = re.compile(r"\(([A-Za-z0-9]+)xx\)?\s*$")
_NUMERIC_RE = re.compile(r"\d+(?:\.\d+)?")

# ---- One-time Oct-26 migration. Remove once this seed has run in
# production (or in the legacy-removal ticket at the latest). ----
#
# Before Oct-26 every Tray node was named with the brand prefix, and the
# MIS/REFLEX ProductFamily nodes were inferred from tray names by regex.
# Renaming the existing Tray nodes in place (rather than letting the seed
# mint new ones) keeps their Procedure REQUIRES, TraySection, and
# BELONGS_TO_TRAY edges. Beyond the prefix strip, three trays were
# genuinely renamed in this catalog; `Pre-Hydrated Bioimplant System |
# Implants` has no successor and goes through the tray-membership reconcile
# like any other dropped tray.
_RENAMED_TRAYS = {
    "Distal Tibia XL Plate Caddy": "Distal Tibia Plating System | XL Plate Caddy",
    "DEX Soft Tissue | Comprehensive Instrument Tray": "Soft Tissue | Comprehensive Instrument Tray",
    "Pre-Hydrated Bioimplant System | Instruments": (
        "Biologics and Bioimplants | Pre-hydrated Bioimplants Instruments"
    ),
}
_RETIRED_FAMILIES = ("MIS", "REFLEX")


def legacy_tray_renames(trays: Iterable[str]) -> dict[str, str]:
    """Old (pre-Oct-26) Tray name -> its name in this catalog, for every
    tray in `trays` that had one."""
    tray_set = set(trays)
    renames = {f"{_BRAND_PREFIX}{tray}": tray for tray in tray_set}
    for old_name, new_name in _RENAMED_TRAYS.items():
        if new_name in tray_set:
            renames[f"{_BRAND_PREFIX}{old_name}"] = new_name
    return renames


# ---- End of one-time migration. ----


class MasterCatalogError(ValueError):
    """The workbook breaks a rule the catalog hierarchy depends on (a
    missing column, a blank Family, a Tray under two Families)."""


@dataclass
class CatalogRow:
    family: str
    tray: str
    item_type: str
    sku: str
    description: str
    qty_per_set: str
    head_style: str
    construct: str
    thread: str
    color: str
    guidewire_spec: str
    pre_drill_spec: str
    driver_spec: str
    indication: str | None
    compat_columns: dict[str, str]


def _clean(value: object) -> str:
    text = "" if value is None else str(value).strip()
    return "" if text.upper() == "N/A" else text


def normalize_tray_name(raw: str) -> str:
    """Drops the leading brand segment (that exact segment only — inner
    " | " parts like "Soft Tissue | Comprehensive Instrument Tray" are part
    of the name) and trims whitespace, which the file doesn't always do."""
    name = raw.strip()
    if name.startswith(_BRAND_PREFIX):
        name = name[len(_BRAND_PREFIX):].strip()
    return name


# 2026-09-04: two independently-run evals of the MIS question set (both
# deterministic and react_agent) found the same defect -- an answer citing
# the wrong thread-type screw for a bunion vs. Akin/fusion procedure. The
# fact needed to get this right (thread <-> indication) already exists in
# this file, just encoded in the description text rather than as its own
# property: MIS Foot Recon's screw descriptions are 100% consistent on this
# (verified directly against the live graph before writing this) --
# "MIS HV CHAMFER" (HV = Hallux Valgus, i.e. bunion) is always Full Thread,
# "MIS COMPRESSION" (compression fixation, the standard for Akin/fusion) is
# always Partial Thread. Scoped to exactly this one tray, not a blanket
# HV/COMPRESSION rule -- both words appear in several *other* trays' rows
# too (plating systems that have nothing to do with MIS), where they almost
# certainly don't carry the same meaning and haven't been verified to.
_MIS_FOOT_RECON_TRAY = "MIS Foot Recon System"


def _infer_indication(tray: str, description: str) -> str | None:
    if tray != _MIS_FOOT_RECON_TRAY:
        return None
    upper = description.upper()
    if "HV" in upper:
        return "Bunion (Hallux Valgus) correction"
    if "COMPRESSION" in upper:
        return "Fusion / Akin osteotomy (compression fixation)"
    return None


def parse_master_catalog(path: Path) -> list[CatalogRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_rows = workbook.worksheets[0].iter_rows(values_only=True)
        headers = [_clean(cell) for cell in next(sheet_rows, ())]
        missing = [column for column in _REQUIRED_COLUMNS if column not in headers]
        if missing:
            raise MasterCatalogError(f"{path.name} has no {', '.join(missing)} column")

        rows = []
        blank_family_rows = []
        families_by_tray: dict[str, set[str]] = defaultdict(set)
        # Row 1 is the header, so data rows start at the spreadsheet's row 2
        # -- the number someone fixing the file would look for.
        for row_number, cells in enumerate(sheet_rows, start=2):
            raw = dict(zip(headers, cells, strict=False))
            sku = _clean(raw.get("Item No."))
            if not sku:
                continue
            family = _clean(raw.get("Family"))
            tray = normalize_tray_name(_clean(raw.get("Tray")))
            if not family:
                blank_family_rows.append(row_number)
                continue
            families_by_tray[tray].add(family)
            compat_columns = {
                header: _clean(value)
                for header, value in raw.items()
                if header and header not in _FIXED_COLUMNS and _clean(value)
            }
            description = _clean(raw.get("Description"))
            rows.append(
                CatalogRow(
                    family=family,
                    tray=tray,
                    item_type=_clean(raw.get("Item Type")),
                    sku=sku,
                    description=description,
                    qty_per_set=_clean(raw.get("Qty per Set")),
                    head_style=_clean(raw.get("Head Style")),
                    construct=_clean(raw.get("Construct")),
                    thread=_clean(raw.get("Thread")),
                    color=_clean(raw.get("Color")),
                    guidewire_spec=_clean(raw.get("Guidewire")),
                    pre_drill_spec=_clean(raw.get("Pre-Drill Diameter")),
                    driver_spec=_clean(raw.get("Driver")),
                    indication=_infer_indication(tray, description),
                    compat_columns=compat_columns,
                )
            )
    finally:
        workbook.close()

    if blank_family_rows:
        raise MasterCatalogError(
            f"{path.name}: every row needs a Family, but row(s) "
            f"{', '.join(str(row_number) for row_number in blank_family_rows)} have none"
        )
    split_trays = {tray: families for tray, families in families_by_tray.items() if len(families) > 1}
    if split_trays:
        details = "; ".join(
            f"{tray!r} under {', '.join(sorted(families))}"
            for tray, families in sorted(split_trays.items())
        )
        raise MasterCatalogError(
            f"{path.name}: each Tray must belong to exactly one Family, but {details}"
        )
    return rows


def _numeric_tokens(text: str) -> set[str]:
    return set(_NUMERIC_RE.findall(text))


def _resolve_by_numeric_tokens(spec: str, keyword: str, candidates: list[CatalogRow]) -> str | None:
    tokens = _numeric_tokens(spec)
    if not tokens:
        return None
    matches = [
        row
        for row in candidates
        if keyword in row.description.upper() and tokens.issubset(_numeric_tokens(row.description))
    ]
    return matches[0].sku if len(matches) == 1 else None


def _resolve_driver_sku(spec: str, candidates: list[CatalogRow]) -> str | None:
    token = spec.strip().upper()
    if not token:
        return None
    pattern = re.compile(rf"\b{re.escape(token)}\b")
    matches = [
        row
        for row in candidates
        if "DRIVER" in row.description.upper() and pattern.search(row.description.upper())
    ]
    return matches[0].sku if len(matches) == 1 else None


def resolve_compatible_skus(header: str, all_skus: list[str], own_sku: str) -> list[str]:
    match = _COMPAT_PREFIX_RE.search(header)
    if not match:
        return []
    prefix = match.group(1)
    return [sku for sku in all_skus if sku != own_sku and sku.startswith(prefix)]


def resolve_tool_sku(kind: str, spec: str, tray_rows: list[CatalogRow]) -> str | None:
    if kind == "guidewire":
        return _resolve_by_numeric_tokens(spec, "GUIDEWIRE", tray_rows)
    if kind == "pre_drill":
        return _resolve_by_numeric_tokens(spec, "DRILL", tray_rows)
    if kind == "driver":
        return _resolve_driver_sku(spec, tray_rows)
    raise ValueError(kind)


async def seed_master_catalog(client: GraphClient, path: Path = DEFAULT_PATH) -> None:
    rows = parse_master_catalog(path)
    if not rows:
        print(f"No rows with an Item No. found in {path}")
        return

    tray_family = {row.tray: row.family for row in rows}
    families = sorted(set(tray_family.values()))

    renamed_trays = await client.rename_trays(legacy_tray_renames(tray_family))

    for family in families:
        await client.upsert_product_family(family)
    for tray, family in tray_family.items():
        await client.upsert_tray(tray, family)

    await client.upsert_parts(
        [
            (
                row.sku,
                row.tray,
                {
                    "description": row.description or None,
                    "item_type": row.item_type or None,
                    "head_style": row.head_style or None,
                    "construct": row.construct or None,
                    "thread": row.thread or None,
                    "color": row.color or None,
                    "guidewire_spec": row.guidewire_spec or None,
                    "pre_drill_spec": row.pre_drill_spec or None,
                    "driver_spec": row.driver_spec or None,
                    "qty_per_set": row.qty_per_set or None,
                    "indication": row.indication,
                },
            )
            for row in rows
        ]
    )

    trays_by_sku: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        trays_by_sku[row.sku].add(row.tray)
    pruned_edges, deleted_trays, kept_trays = await client.prune_stale_tray_memberships(
        {sku: sorted(trays) for sku, trays in trays_by_sku.items()}
    )
    retired_families = await client.delete_product_families(
        [family for family in _RETIRED_FAMILIES if family not in families]
    )

    all_skus = [row.sku for row in rows]
    compatible_pairs: list[tuple[str, str]] = []
    for row in rows:
        for header, value in row.compat_columns.items():
            if value.upper() != "X":
                continue
            for other_sku in resolve_compatible_skus(header, all_skus, row.sku):
                compatible_pairs.append((row.sku, other_sku))
    await client.upsert_compatible_with_pairs(compatible_pairs)
    compat_count = len(compatible_pairs)

    rows_by_tray: dict[str, list[CatalogRow]] = defaultdict(list)
    for row in rows:
        rows_by_tray[row.tray].append(row)

    tool_pairs: list[tuple[str, str]] = []
    for row in rows:
        tray_rows = rows_by_tray[row.tray]
        for kind, spec in (
            ("guidewire", row.guidewire_spec),
            ("pre_drill", row.pre_drill_spec),
            ("driver", row.driver_spec),
        ):
            if not spec:
                continue
            tool_sku = resolve_tool_sku(kind, spec, tray_rows)
            if tool_sku and tool_sku != row.sku:
                tool_pairs.append((row.sku, tool_sku))
    await client.upsert_requires_tool_pairs(tool_pairs)
    tool_count = len(tool_pairs)

    trays_by_family: dict[str, list[str]] = defaultdict(list)
    for tray, family in sorted(tray_family.items()):
        trays_by_family[family].append(tray)
    await upsert_families_and_trays(dict(trays_by_family))

    print(
        f"Seeded {len(rows)} parts across {len(families)} product families and "
        f"{len(tray_family)} trays (mirrored to Postgres), "
        f"{compat_count} COMPATIBLE_WITH edges, {tool_count} REQUIRES_TOOL edges; "
        f"removed {pruned_edges} stale BELONGS_TO_TRAY edges"
    )
    for old_name, new_name in renamed_trays:
        print(f"  renamed tray: {old_name} -> {new_name}")
    for family in retired_families:
        print(f"  retired product family: {family}")
    for tray in deleted_trays:
        print(f"  deleted emptied tray: {tray}")
    for tray in kept_trays:
        print(f"  kept emptied tray (still referenced, e.g. by a Procedure): {tray}")


async def main() -> None:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_PATH
    client = get_graph_client()
    await client.ensure_constraints()
    await seed_master_catalog(client, path)


if __name__ == "__main__":
    asyncio.run(main())
